"""Extract pair rows from the five MVP2 seed supplements and screen new coverage.

The script writes a source-traceable candidate CSV and an anti-join result. It
does not decide whether a reported association is scientifically valid.
"""

from __future__ import annotations

import argparse
import csv
import json
import math
from pathlib import Path
import sys
from datetime import datetime, timezone
from typing import Any

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
SEEDS = ROOT / "MVP2" / "literature" / "seeds"
BASELINE = ROOT / "archive" / "project_files" / "release-evidence_2026-09-11" / "metabocausal-public" / "data" / "metabolite_disease.csv"
sys.path.insert(0, str(ROOT / "MVP2"))
from src.pair_screening import hashed, read_csv, screen_pairs  # noqa: E402


_EMPTY_CELL_VALUES = {"", "none", "nan", "null", "n/a", "na", "not available", "-"}


def text(row: tuple[Any, ...], index: int) -> str:
    """Return a source cell as text, treating blank Excel cells as blank.

    ``str(None)`` is the literal string ``"None"``.  The previous extractor
    accidentally turned footer and blank rows into fake metabolite/disease
    records because of that conversion.
    """
    value = row[index] if index < len(row) else None
    if value is None:
        return ""
    value = " ".join(str(value).split())
    return "" if value.casefold() in _EMPTY_CELL_VALUES else value


def _as_float(value: str) -> float | None:
    """The numeric value of an extracted cell, or ``None`` when it is not a number."""
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


# Written onto a row whose two 95% bounds came out of the source in the impossible order.
# The reorder is a per-row numeric decision, never a blanket re-reading of the header: a
# row whose bounds already ascend is written exactly as it was read.
_REORDERED_BOUNDS_NOTE = (
    "95% bounds reordered: read under the source header's own names this row came out "
    "with ci_low greater than ci_high, which no confidence interval can satisfy, so the "
    "two numbers are written to ci_low/ci_high in numeric order. Neither number was "
    "recomputed - the pair is the source's own - and the values exactly as the source "
    "prints them are preserved in ci_low_as_printed/ci_high_as_printed. Observed on the "
    "UKB mapping supplement, whose header prints '95% UCI' over the column holding the "
    "lower bound and '95% LCI' over the column holding the upper one."
)


def bounds_in_numeric_order(ci_low: str, ci_high: str) -> tuple[str, str, bool]:
    """The two 95% bounds in numeric order, and whether they had to be swapped.

    ``ci_low > ci_high`` is not a judgement about what a source meant; it is an
    arithmetic impossibility, and it is the only trigger here. A row whose bounds
    already ascend, whose bounds are equal, or either of whose bounds is blank or
    non-numeric is returned untouched, so nothing is reordered by reading the
    printed header rather than the printed values.
    """
    low_value, high_value = _as_float(ci_low), _as_float(ci_high)
    if low_value is None or high_value is None or low_value <= high_value:
        return ci_low, ci_high, False
    return ci_high, ci_low, True


def add(records: list[dict[str, Any]], *, study_id: str, path: Path, sheet: str, row_number: int,
        metabolite_name: str, disease_trait: str, direction: str, metabolite_id: str = "",
        disease_id: str = "", method: str = "", beta: str = "", se: str = "",
        p_value: str = "", ci_low: str = "", ci_high: str = "", nsnp: str = "",
        ci_low_beta_scale: str = "", ci_high_beta_scale: str = "",
        extraction_note: str = "") -> None:
    if not metabolite_name or not disease_trait:
        return
    ordered_low, ordered_high, reordered = bounds_in_numeric_order(ci_low, ci_high)
    ci_low_as_printed = ci_low if reordered else ""
    ci_high_as_printed = ci_high if reordered else ""
    if reordered:
        extraction_note = " | ".join(
            note for note in (extraction_note, _REORDERED_BOUNDS_NOTE) if note
        )
    ci_low, ci_high = ordered_low, ordered_high
    records.append({
        "study_id": study_id, "source_path": path.relative_to(ROOT).as_posix(),
        "source_sheet": sheet, "source_row": row_number,
        "metabolite_name": metabolite_name, "metabolite_id": metabolite_id,
        "disease_trait": disease_trait, "disease_id": disease_id,
        "direction": direction, "method": method, "beta": beta, "se": se,
        "p_value": p_value, "ci_low": ci_low, "ci_high": ci_high, "nsnp": nsnp,
        "ci_low_beta_scale": ci_low_beta_scale, "ci_high_beta_scale": ci_high_beta_scale,
        "ci_low_as_printed": ci_low_as_printed, "ci_high_as_printed": ci_high_as_printed,
        "extraction_note": extraction_note,
        "pair_relation_verified": "yes",
        "pair_relation_basis": "same source row has an explicit metabolite/exposure field and disease/outcome field",
        "record_type": "paper_screening_candidate",
    })


def extract_chen(records: list[dict[str, Any]]) -> None:
    path = SEEDS / "chen_2023" / "41588_2022_1270_MOESM4_ESM.xlsx"
    wb = load_workbook(path, read_only=True, data_only=True)
    for sheet in ("ST18", "ST19", "ST20"):
        for row_number, row in enumerate(wb[sheet].iter_rows(min_row=4, values_only=True), 4):
            add(records, study_id="chen_2023", path=path, sheet=sheet, row_number=row_number,
                metabolite_name=text(row, 0), disease_trait=text(row, 1), direction="forward",
                method=text(row, 4), nsnp=text(row, 5), beta=text(row, 6), ci_low=text(row, 7),
                ci_high=text(row, 8), p_value=text(row, 9))
    for row_number, row in enumerate(wb["ST23"].iter_rows(min_row=4, values_only=True), 4):
        add(records, study_id="chen_2023", path=path, sheet="ST23", row_number=row_number,
            metabolite_name=text(row, 2), disease_trait=text(row, 1), direction="reverse",
            method=text(row, 3), nsnp=text(row, 4), beta=text(row, 5), se=text(row, 6), p_value=text(row, 7))


def extract_yin(records: list[dict[str, Any]]) -> None:
    path = SEEDS / "yin_2025" / "41467_2025_58129_MOESM4_ESM.xlsx"
    ws = load_workbook(path, read_only=True, data_only=True)["Sheet1"]
    for row_number, row in enumerate(ws.iter_rows(min_row=8, values_only=True), 8):
        add(records, study_id="yin_2025", path=path, sheet="Sheet1", row_number=row_number,
            metabolite_id=text(row, 0), metabolite_name=text(row, 1), disease_id=text(row, 3),
            disease_trait=text(row, 4), direction="forward", method="MR-RAPS", beta=text(row, 13),
            se=text(row, 14), p_value=text(row, 15))


def extract_zoodsma(records: list[dict[str, Any]]) -> None:
    path = SEEDS / "zoodsma_2025" / "supplements" / "41588_2025_2355_MOESM4_ESM.xlsx"
    ws = load_workbook(path, read_only=True, data_only=True)["ST14"]
    for row_number, row in enumerate(ws.iter_rows(min_row=4, values_only=True), 4):
        add(records, study_id="zoodsma_2025", path=path, sheet="ST14", row_number=row_number,
            metabolite_name=text(row, 1), disease_trait=text(row, 3), direction="forward",
            method="Inverse-Variance-Weighted 2-Sample MR", nsnp=text(row, 5), beta=text(row, 6),
            se=text(row, 7), p_value=text(row, 8))


_ODDS_RATIO_RELATIVE_TOLERANCE = 1e-6
_ODDS_RATIO_PRINTED_PRECISION_TOLERANCE = 0.005


def _is_number(value: Any) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def _matches_odds_ratio(value: Any, expected: float) -> bool:
    if not _is_number(value):
        return False
    if expected == 0:
        return abs(value) < _ODDS_RATIO_RELATIVE_TOLERANCE
    return abs(value - expected) / abs(expected) < _ODDS_RATIO_RELATIVE_TOLERANCE


def ukb_row_layout_is_shifted(row: tuple[Any, ...]) -> bool:
    """True when the printed header's column order does not hold for this row.

    sTables 31-35 print one 'Beta, Odds ratio, 95% UCI, 95% LCI, p value'
    header over columns 9-13 for every row, but a single-SNP 'Wald ratio' row
    carries its values one column to the right of that header. Reading the
    fixed indices on such a row puts the odds ratio in ci_high and a
    beta-scale bound in p_value. exp(beta) is an identity the row must
    satisfy either way, so it names which physical column holds the odds
    ratio independently of the printed header.
    """
    beta = row[9] if len(row) > 9 else None
    if not _is_number(beta):
        return False
    try:
        expected_odds_ratio = math.exp(beta)
    except OverflowError:
        return False
    header_column = row[10] if len(row) > 10 else None
    if _matches_odds_ratio(header_column, expected_odds_ratio):
        return False
    if _is_number(header_column) and abs(header_column - expected_odds_ratio) <= _ODDS_RATIO_PRINTED_PRECISION_TOLERANCE:
        return False
    shifted_column = row[11] if len(row) > 11 else None
    return _matches_odds_ratio(shifted_column, expected_odds_ratio)


# What the shifted rows' own 95% bounds are, and why they are not written into
# ci_low/ci_high. On an unshifted row those two columns are on the ODDS-RATIO scale; on a
# shifted Wald-ratio row the same physical columns hold the interval on the BETA scale.
# Writing them into the same field would mix two scales in one column, so they are carried
# in their own columns and the row says what happened. Dropping them silently is the one
# option this project's rules exclude.
_SHIFTED_ROW_NOTE = (
    "single-SNP Wald-ratio row: the printed 'Beta, Odds ratio, 95% UCI, 95% LCI, p value' "
    "header does not hold for this row. The odds ratio is one column right of the header, "
    "so the p value was read from column 10 and the 95% bounds, which are on the beta "
    "scale here, are carried in ci_low_beta_scale/ci_high_beta_scale rather than in "
    "ci_low/ci_high, which hold odds-ratio-scale bounds on every other row."
)

# On an unshifted UKB row the printed header labels column 11 '95% UCI' and column 12
# '95% LCI'. The values are the other way round: column 11 holds the LOWER bound. Reading
# the two cells under the header's own names therefore produced ci_low > ci_high, which no
# confidence interval can satisfy. The shipped ci_low/ci_high now hold the two numbers in
# numeric order; the values as the source prints them are preserved on the same row in
# ci_low_as_printed/ci_high_as_printed, so nothing the source published is lost and a
# reader can see both what was printed and what was shipped.
_BOUNDS_ORDER_NOTE = (
    "UKB rows: the printed header names column 11 '95% UCI' and column 12 '95% LCI', and "
    "the values are the other way round - column 11 holds the lower bound. Reading them "
    "under the header's own names yields ci_low > ci_high, which is impossible for a "
    "confidence interval. ci_low/ci_high therefore hold the two source numbers in numeric "
    "order, every reordered row says so in extraction_note, and the as-printed values are "
    "preserved in ci_low_as_printed/ci_high_as_printed. No bound was recomputed and no row "
    "whose bounds already ascend was touched."
)


def extract_ukb(records: list[dict[str, Any]]) -> None:
    path = SEEDS / "ukb_mapping_2025" / "42255_2025_1371_MOESM3_ESM.xlsx"
    wb = load_workbook(path, read_only=True, data_only=True)
    for sheet in ("sTable 31", "sTable 32", "sTable 33", "sTable 34", "sTable 35"):
        direction = "reverse" if sheet == "sTable 35" else "forward"
        for row_number, row in enumerate(wb[sheet].iter_rows(min_row=4, values_only=True), 4):
            if ukb_row_layout_is_shifted(row):
                ci_high, ci_low, p_value = "", "", text(row, 10)
                ci_low_beta, ci_high_beta = text(row, 12), text(row, 13)
                note = _SHIFTED_ROW_NOTE
            else:
                ci_high, ci_low, p_value = text(row, 11), text(row, 12), text(row, 13)
                ci_low_beta, ci_high_beta = "", ""
                note = ""
            add(
                records,
                study_id          = "ukb_mapping_2025",
                path              = path,
                sheet             = sheet,
                row_number        = row_number,
                metabolite_id     = text(row, 0),
                metabolite_name   = text(row, 1),
                disease_id        = text(row, 3),
                disease_trait     = text(row, 4),
                direction         = direction,
                method            = text(row, 8),
                nsnp              = text(row, 7),
                beta              = text(row, 9),
                ci_high           = ci_high,
                ci_low            = ci_low,
                p_value           = p_value,
                ci_low_beta_scale = ci_low_beta,
                ci_high_beta_scale = ci_high_beta,
                extraction_note   = note,
            )


def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    fields: list[str] = []
    for row in rows:
        for field in row:
            if field not in fields:
                fields.append(field)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--baseline", type=Path, default=BASELINE)
    parser.add_argument("--output-dir", type=Path, default=ROOT / "MVP2" / "data" / "extracted" / "screened_seed_papers")
    args = parser.parse_args()
    records: list[dict[str, Any]] = []
    extract_chen(records)
    extract_yin(records)
    extract_zoodsma(records)
    extract_ukb(records)
    screened = screen_pairs(read_csv(args.baseline), records)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    candidates_path = args.output_dir / "candidate_pairs.csv"
    new_path = args.output_dir / "new_pair_records.csv"
    unique_path = args.output_dir / "new_unique_pairs.csv"
    ledger_path = args.output_dir / "screening_ledger.csv"
    write_csv(candidates_path, records)
    write_csv(new_path, screened["new_pair_records"])
    write_csv(unique_path, screened["new_unique_pairs"])
    write_csv(ledger_path, screened["decisions"])
    shifted_rows = sum(1 for row in records if _SHIFTED_ROW_NOTE in row["extraction_note"])
    reordered_rows = sum(1 for row in records if _REORDERED_BOUNDS_NOTE in row["extraction_note"])
    reordered_new_pair_rows = sum(
        1 for row in screened["new_pair_records"] if _REORDERED_BOUNDS_NOTE in row["extraction_note"]
    )
    reordered_unique_pair_rows = sum(
        1 for row in screened["new_unique_pairs"] if _REORDERED_BOUNDS_NOTE in row["extraction_note"]
    )
    manifest = {
        "schema_version": 2,
        "run_id": datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ"),
        "run_id_note": (
            "run_id is when these files were last written. The directory name records the "
            "screening these tables belong to; re-running the same extractor over the same "
            "seed workbooks and the same baseline rewrites the files without making it a "
            "different screening."
        ),
        "baseline": hashed(args.baseline),
        "candidate_sources": [
            "Chen 2023 ST18-ST20 forward and ST23 reverse",
            "Yin 2025 Supplementary Data 2 forward",
            "Zoodsma 2025 ST14 forward",
            "UKB Mapping 2025 sTables 31-34 forward and 35 reverse",
            "Shin 2014: no metabolite-disease MR table extracted",
        ],
        "outputs": {"candidate_pairs": candidates_path.name, "new_pair_records": new_path.name, "new_unique_pairs": unique_path.name, "screening_ledger": ledger_path.name},
        "summary": screened["summary"],
        "bound_order": {
            "rule": (
                "A row is reordered only when its two extracted 95% bounds are both "
                "numeric and ci_low > ci_high, which no confidence interval can satisfy. "
                "Rows already in ascending order, rows with equal bounds and rows with a "
                "blank or non-numeric bound are written exactly as they were read."
            ),
            "reordered_candidate_rows": reordered_rows,
            "reordered_new_pair_rows": reordered_new_pair_rows,
            "reordered_unique_pair_rows": reordered_unique_pair_rows,
            "as_printed_columns": ["ci_low_as_printed", "ci_high_as_printed"],
            "as_printed_note": (
                "Populated only on a reordered row, where they hold that row's bounds "
                "exactly as the source prints them; blank on every row that was not "
                "reordered."
            ),
            "remaining_ci_low_greater_than_ci_high": sum(
                1
                for row in records
                if (low := _as_float(row["ci_low"])) is not None
                and (high := _as_float(row["ci_high"])) is not None
                and low > high
            ),
        },
        "extraction_rule": "Keep a row only when the Excel row itself contains both a non-empty metabolite/exposure value and a non-empty disease/outcome value. Blank, footer, header, and placeholder cells are excluded; no pair is inferred from a paper title or from separate rows.",
        "caveats": [
            "This is source extraction plus an anti-join, not scientific validation or acceptance of MR evidence.",
            "Reverse and unresolved rows remain in the ledger but are excluded from new_pair_records.csv.",
            "Phenotype terminology is preserved as printed in the source; disease/outcome interpretation remains a screening decision.",
            f"{shifted_rows} extracted UKB row(s) carry a shifted source layout. {_SHIFTED_ROW_NOTE}",
            f"{reordered_rows} extracted row(s) had their two 95% bounds reordered into numeric order. {_BOUNDS_ORDER_NOTE}",
            "Pair identity keeps '%': a ratio exposure such as 'Omega-3%' and the absolute exposure 'Omega-3' are two different exposures and are two different pairs.",
        ],
    }
    (args.output_dir / "manifest.json").write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(manifest["summary"]))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
