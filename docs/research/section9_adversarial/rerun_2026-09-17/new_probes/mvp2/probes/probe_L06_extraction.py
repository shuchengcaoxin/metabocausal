"""L06 - does the fixed-column extractor agree with the source workbooks?

L05 found that most kept rows carry ci_low > ci_high and that seven carry a p_value
outside [0, 1]. This probe goes back to the committed seed workbooks and asks which
column actually holds which quantity, using an arithmetic identity the source cannot
fake: for a log-odds effect the odds-ratio column must equal exp(beta).

  1. per sheet, which column index satisfies  col ~= exp(Beta)  - this labels the row's
     true layout independently of the sheet's printed header;
  2. what MVP2/scripts/screen_seed_papers.py reads from each fixed index;
  3. how many extracted rows therefore carry a value that is not the quantity its column
     is named for, and how many of those survive into new_unique_pairs.csv.

Deterministic, offline, read-only. Nothing under MVP2/ is written.

Run from the repository root:
    python docs/research/section9_adversarial/rerun_2026-09-17/new_probes/mvp2/probes/probe_L06_extraction.py
"""

from __future__ import annotations

import io
import json
import math
import os
import sys
from collections import Counter, defaultdict


def _find_repo_root(start: str) -> str:
    here = os.path.abspath(start)
    while True:
        if (
            os.path.isfile(os.path.join(here, "render.py"))
            and os.path.isfile(os.path.join(here, "classify.py"))
            and os.path.isdir(os.path.join(here, "MVP2"))
        ):
            return here
        parent = os.path.dirname(here)
        if parent == here:
            raise RuntimeError("repository root not found above %s" % start)
        here = parent


HERE = os.path.dirname(os.path.abspath(__file__))
REPO = _find_repo_root(HERE)
OUT_DIR = os.path.abspath(os.path.join(HERE, "..", "outputs"))
sys.path.insert(0, os.path.join(REPO, "MVP2"))

from openpyxl import load_workbook  # noqa: E402
from src.pair_screening import read_csv  # noqa: E402

UKB = os.path.join(
    REPO, "MVP2", "literature", "seeds", "ukb_mapping_2025", "42255_2025_1371_MOESM3_ESM.xlsx"
)
SHEETS = ["sTable 31", "sTable 32", "sTable 33", "sTable 34", "sTable 35"]
SCREENED = os.path.join(REPO, "MVP2", "data", "extracted", "screened_20260912")

# MVP2/scripts/screen_seed_papers.py:106 - the fixed indices the extractor reads.
EXTRACTOR_MAP = {
    "metabolite_id": 0,
    "metabolite_name": 1,
    "disease_id": 3,
    "disease_trait": 4,
    "nsnp": 7,
    "method": 8,
    "beta": 9,
    "ci_high": 11,
    "ci_low": 12,
    "p_value": 13,
}
TOLERANCE = 1e-6


def number(value):
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    return None


def close(a, b):
    if a is None or b is None:
        return False
    if b == 0:
        return abs(a) < TOLERANCE
    return abs(a - b) / abs(b) < TOLERANCE


def main() -> int:
    os.makedirs(OUT_DIR, exist_ok=True)
    lines = ["L06 - fixed-column extraction vs the ukb_mapping_2025 source workbook", ""]
    lines.append("source: MVP2/literature/seeds/ukb_mapping_2025/42255_2025_1371_MOESM3_ESM.xlsx")
    lines.append("extractor: MVP2/scripts/screen_seed_papers.py:extract_ukb, fixed indices")
    lines.append("           %s" % json.dumps(EXTRACTOR_MAP, sort_keys=True))
    lines.append("")
    record = {"extractor_map": EXTRACTOR_MAP}

    workbook = load_workbook(UKB, read_only=True, data_only=True)

    # Printed header, for the record.
    lines.append("[0] printed header of each sheet (row 3, the sub-header the values sit under)")
    headers = {}
    for sheet in SHEETS:
        rows = workbook[sheet].iter_rows(min_row=3, max_row=3, values_only=True)
        header = list(next(rows))[:15]
        headers[sheet] = header
        lines.append("  %-12s %s" % (sheet, header[9:15]))
    record["printed_headers"] = {k: [str(v) for v in val] for k, val in headers.items()}

    # ------------------------------------------------------------------ 1. layout test
    lines.append("")
    lines.append("[1] which column satisfies  col ~= exp(Beta at col 9)  (relative tol %g)" % TOLERANCE)
    lines.append("    col10 == exp(beta) -> the printed header order holds ('layout H')")
    lines.append("    col11 == exp(beta) -> the row is shifted one column right ('layout S')")
    lines.append("")
    layouts = defaultdict(Counter)
    shifted_rows = defaultdict(list)
    header_rows = defaultdict(list)
    for sheet in SHEETS:
        for row_number, row in enumerate(
            workbook[sheet].iter_rows(min_row=4, values_only=True), 4
        ):
            beta = number(row[9] if len(row) > 9 else None)
            if beta is None:
                layouts[sheet]["no numeric beta"] += 1
                continue
            try:
                expected = math.exp(beta)
            except OverflowError:
                layouts[sheet]["beta overflows exp()"] += 1
                continue
            c10 = number(row[10] if len(row) > 10 else None)
            c11 = number(row[11] if len(row) > 11 else None)
            if close(c10, expected):
                layouts[sheet]["layout H (col10 = OR)"] += 1
                header_rows[sheet].append(row_number)
            elif close(c11, expected):
                layouts[sheet]["layout S (col11 = OR)"] += 1
                shifted_rows[sheet].append(row_number)
            elif c10 is not None and abs(c10 - expected) <= 0.005:
                # sTable 34 prints the odds ratio to two decimals; the identity still
                # holds within the printed precision, so this is layout H, not a shift.
                layouts[sheet]["layout H (col10 = OR, printed to 2 decimals)"] += 1
                header_rows[sheet].append(row_number)
            else:
                layouts[sheet]["neither col10 nor col11 = exp(beta)"] += 1
    for sheet in SHEETS:
        lines.append("  %s" % sheet)
        for label in sorted(layouts[sheet]):
            lines.append("    %-38s %d" % (label, layouts[sheet][label]))
    record["layout_counts"] = {k: dict(v) for k, v in layouts.items()}

    # ------------------------------------------------------------------ 2. what that means
    lines.append("")
    lines.append("[2] what the extractor writes for a 'layout S' row")
    example_sheet = next((s for s in SHEETS if shifted_rows[s]), None)
    if example_sheet:
        example_row_no = shifted_rows[example_sheet][0]
        for row_number, row in enumerate(
            workbook[example_sheet].iter_rows(min_row=4, values_only=True), 4
        ):
            if row_number == example_row_no:
                beta = number(row[9])
                lines.append("  %s row %d" % (example_sheet, row_number))
                for index in (9, 10, 11, 12, 13, 14):
                    label = headers[example_sheet][index] if index < len(headers[example_sheet]) else "?"
                    lines.append("    col%-3d header %-14r value %r" % (index, label, row[index]))
                lines.append("    exp(col9) = %.12g   (matches col11, not col10)" % math.exp(beta))
                lines.append("    col12 and col13 bracket col9 symmetrically: %s < %s < %s"
                             % (row[12], row[9], row[13]))
                lines.append("")
                lines.append("    the extractor therefore records")
                lines.append("      beta    <- col9  = %r   (the beta)" % (row[9],))
                lines.append("      ci_high <- col11 = %r   (the odds ratio)" % (row[11],))
                lines.append("      ci_low  <- col12 = %r   (a lower bound on the beta scale)" % (row[12],))
                lines.append("      p_value <- col13 = %r   (an upper bound, not a p-value)" % (row[13],))
                record["layout_S_example"] = {
                    "sheet": example_sheet,
                    "row": row_number,
                    "cols_9_to_14": [str(row[i]) for i in range(9, 15)],
                }
                break

    # ------------------------------------------------------------------ 3. survivors
    lines.append("")
    lines.append("[3] how many affected rows reach the kept tables")
    lines.append("")
    shifted_index = {(sheet, str(r)) for sheet in SHEETS for r in shifted_rows[sheet]}
    tables = {
        "candidate_pairs.csv": read_csv(os.path.join(SCREENED, "candidate_pairs.csv")),
        "new_pair_records.csv": read_csv(os.path.join(SCREENED, "new_pair_records.csv")),
        "new_unique_pairs.csv": read_csv(os.path.join(SCREENED, "new_unique_pairs.csv")),
    }
    survivor_counts = {}
    for name in sorted(tables):
        rows = tables[name]
        ukb_rows = [r for r in rows if r["study_id"] == "ukb_mapping_2025"]
        affected = [
            r for r in ukb_rows
            if (r["source_sheet"], str(r["source_row"])) in shifted_index
        ]
        p_outside = 0
        for row in affected:
            try:
                value = float(row["p_value"])
            except (TypeError, ValueError):
                continue
            if value < 0 or value > 1:
                p_outside += 1
        survivor_counts[name] = {
            "rows_total": len(rows),
            "rows_from_ukb_mapping_2025": len(ukb_rows),
            "rows_on_layout_S": len(affected),
            "of_those_p_value_outside_0_1": p_outside,
            "of_those_p_value_inside_0_1_and_therefore_undetectable_by_range": len(affected) - p_outside,
        }
        lines.append("  %s" % name)
        for key in (
            "rows_total",
            "rows_from_ukb_mapping_2025",
            "rows_on_layout_S",
            "of_those_p_value_outside_0_1",
            "of_those_p_value_inside_0_1_and_therefore_undetectable_by_range",
        ):
            lines.append("    %-62s %d" % (key, survivor_counts[name][key]))
    record["survivors"] = survivor_counts

    # ------------------------------------------------------------------ 4. CI ordering
    lines.append("")
    lines.append("[4] confidence-bound ordering on 'layout H' rows (the printed-header case)")
    lines.append("    header says col11 = '95% UCI', col12 = '95% LCI';")
    lines.append("    extractor maps ci_high <- col11 and ci_low <- col12.")
    lines.append("")
    ordering = {}
    for sheet in SHEETS:
        wrong = right = missing = 0
        for row_number, row in enumerate(
            workbook[sheet].iter_rows(min_row=4, values_only=True), 4
        ):
            if row_number not in set(header_rows[sheet]):
                continue
            lower_named_upper = number(row[11] if len(row) > 11 else None)
            upper_named_lower = number(row[12] if len(row) > 12 else None)
            if lower_named_upper is None or upper_named_lower is None:
                missing += 1
            elif lower_named_upper > upper_named_lower:
                right += 1  # value under '95% UCI' really is the larger one
            else:
                wrong += 1  # value under '95% UCI' is the SMALLER one
        ordering[sheet] = {
            "layout_H_rows": len(header_rows[sheet]),
            "col11_is_larger_than_col12": right,
            "col11_is_smaller_than_col12": wrong,
            "a_bound_missing": missing,
        }
        lines.append(
            "  %-12s layout-H rows %-6d  col11>col12 %-6d  col11<col12 %-6d  missing %d"
            % (
                sheet,
                ordering[sheet]["layout_H_rows"],
                right,
                wrong,
                missing,
            )
        )
    lines.append("")
    lines.append(
        "  A row where the value printed under '95% UCI' is smaller than the value under"
    )
    lines.append(
        "  '95% LCI' means the source's two bound headers do not describe the numbers under"
    )
    lines.append(
        "  them; the extractor follows the headers, so ci_low and ci_high come out inverted."
    )
    record["ci_ordering"] = ordering

    # the other three seed extractions, for contrast
    lines.append("")
    lines.append("[5] contrast: bound ordering in the kept table, by study and sheet")
    kept = tables["new_unique_pairs.csv"]
    contrast = defaultdict(lambda: {"ci_low>ci_high": 0, "ci_low<=ci_high": 0, "no bounds": 0})
    for row in kept:
        key = "%s / %s" % (row["study_id"], row["source_sheet"])
        try:
            low = float(row["ci_low"])
            high = float(row["ci_high"])
        except (TypeError, ValueError):
            contrast[key]["no bounds"] += 1
            continue
        contrast[key]["ci_low>ci_high" if low > high else "ci_low<=ci_high"] += 1
    for key in sorted(contrast):
        entry = contrast[key]
        lines.append(
            "  %-34s inverted %-6d ordered %-6d no bounds %d"
            % (key, entry["ci_low>ci_high"], entry["ci_low<=ci_high"], entry["no bounds"])
        )
    record["kept_bound_ordering"] = {k: dict(v) for k, v in contrast.items()}

    # ------------------------------------------------------------------ 6. sheet survival
    lines.append("")
    lines.append("[6] rows per source sheet at each stage (ukb_mapping_2025 only)")
    lines.append("")
    lines.append(
        "  %-12s %-18s %-20s %s"
        % ("sheet", "candidate_pairs", "new_pair_records", "new_unique_pairs")
    )
    survival = {}
    sheets_seen = sorted(
        {r["source_sheet"] for r in tables["candidate_pairs.csv"] if r["study_id"] == "ukb_mapping_2025"}
    )
    for sheet in sheets_seen:
        counts = []
        for name in ("candidate_pairs.csv", "new_pair_records.csv", "new_unique_pairs.csv"):
            counts.append(
                sum(
                    1
                    for r in tables[name]
                    if r["study_id"] == "ukb_mapping_2025" and r["source_sheet"] == sheet
                )
            )
        survival[sheet] = dict(
            zip(
                ("candidate_pairs", "new_pair_records", "new_unique_pairs"),
                counts,
            )
        )
        lines.append("  %-12s %-18d %-20d %d" % (sheet, counts[0], counts[1], counts[2]))
    record["sheet_survival"] = survival

    text = "\n".join(lines) + "\n"
    with io.open(os.path.join(OUT_DIR, "L06_extraction.txt"), "w", encoding="utf-8") as handle:
        handle.write(text)
    with io.open(os.path.join(OUT_DIR, "L06_extraction.json"), "w", encoding="utf-8") as handle:
        json.dump(record, handle, indent=2, sort_keys=True, default=str)
        handle.write("\n")
    sys.stdout.write(text)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
